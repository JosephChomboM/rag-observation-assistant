# Databricks notebook source
# MAGIC %pip install openai python-docx openpyxl beautifulsoup4 pillow

# COMMAND ----------

"""
Databricks METOD RAG Observation Generator MVP.

This notebook validates the end-to-end flow:
stable METOD knowledge + runtime evidence + model context -> structured observation.

It intentionally avoids Delta tables, Spark DataFrames, web apps, external vector
databases, LangChain, and autonomous-agent patterns.
"""

# COMMAND ----------

import base64
import json
import math
import mimetypes
import os
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from bs4 import BeautifulSoup
from docx import Document
from openai import AzureOpenAI
from openpyxl import load_workbook

# COMMAND ----------

# ============================================================
# AZURE OPENAI CONFIGURATION
# ============================================================

AZURE_OPENAI_ENDPOINT = os.getenv("AZURE_OPENAI_ENDPOINT", "").strip()
AZURE_OPENAI_API_KEY = os.getenv("AZURE_OPENAI_API_KEY", "").strip()
AZURE_OPENAI_API_VERSION = os.getenv("AZURE_OPENAI_API_VERSION", "2024-10-21").strip()
EMBEDDING_DEPLOYMENT = os.getenv(
    "AZURE_OPENAI_EMBEDDING_DEPLOYMENT",
    "text-embedding-3-small",
).strip()
CHAT_DEPLOYMENT = os.getenv("AZURE_OPENAI_CHAT_DEPLOYMENT", "gpt-4o-mini").strip()


def require_azure_openai_credentials() -> None:
    missing = []
    if not AZURE_OPENAI_ENDPOINT:
        missing.append("AZURE_OPENAI_ENDPOINT")
    if not AZURE_OPENAI_API_KEY:
        missing.append("AZURE_OPENAI_API_KEY")
    if missing:
        raise ValueError(
            "Missing Azure OpenAI environment variable(s): "
            + ", ".join(missing)
            + ". Configure them on the Databricks cluster before running the notebook."
        )


require_azure_openai_credentials()

azure_client = AzureOpenAI(
    azure_endpoint=AZURE_OPENAI_ENDPOINT,
    api_key=AZURE_OPENAI_API_KEY,
    api_version=AZURE_OPENAI_API_VERSION,
)

print("Azure OpenAI endpoint configured.")
print("Embedding deployment:", EMBEDDING_DEPLOYMENT)
print("Chat deployment:", CHAT_DEPLOYMENT)

# COMMAND ----------

# ============================================================
# DATABRICKS WIDGETS
# ============================================================

METOD_CHOICES = [
    "METOD1",
    "METOD2",
    "METOD3",
    "METOD4",
    "METOD5",
    "METOD6",
    "METOD7",
    "METOD8",
]

dbutils.widgets.dropdown(
    name="metod_id",
    defaultValue="METOD1",
    choices=METOD_CHOICES,
    label="METOD folder",
)
dbutils.widgets.dropdown(
    name="rebuild_index",
    defaultValue="true",
    choices=["true", "false"],
    label="Rebuild index",
)
dbutils.widgets.text("model_name", "", "Model name")
dbutils.widgets.text("model_type", "Scoring model", "Model type")
dbutils.widgets.text("business_objective", "", "Business objective")
dbutils.widgets.text("population", "", "Population")
dbutils.widgets.text("target_variable", "", "Target variable")
dbutils.widgets.text("additional_notes", "", "Additional notes")

METOD_ID = dbutils.widgets.get("metod_id").strip().upper()
REBUILD_INDEX = dbutils.widgets.get("rebuild_index").strip().lower() == "true"

if METOD_ID not in METOD_CHOICES:
    raise ValueError(f"Unsupported METOD '{METOD_ID}'. Expected one of: {METOD_CHOICES}")

MODEL_CONTEXT = {
    "model_name": dbutils.widgets.get("model_name").strip(),
    "model_type": dbutils.widgets.get("model_type").strip(),
    "business_objective": dbutils.widgets.get("business_objective").strip(),
    "population": dbutils.widgets.get("population").strip(),
    "target_variable": dbutils.widgets.get("target_variable").strip(),
    "additional_notes": dbutils.widgets.get("additional_notes").strip(),
}

print("Selected METOD:", METOD_ID)
print("Rebuild index:", REBUILD_INDEX)
print("Model context:", json.dumps(MODEL_CONTEXT, indent=2, ensure_ascii=False))

# COMMAND ----------

# ============================================================
# STORAGE CONFIGURATION
# ============================================================

USER_EMAIL = "josephchombo@company.com"

# Runtime artifacts live in Workspace Files. Use normal Python file operations.
WORKSPACE_BASE_DIR = f"/Workspace/Users/{USER_EMAIL}/softia_mvp"

# Use this only if a Databricks API requires file:/ paths.
WORKSPACE_BASE_URI = f"file:{WORKSPACE_BASE_DIR}"

# Stable METOD knowledge lives in FileStore or Blob/ADLS.
KNOWLEDGE_BASE_DIR = "/dbfs/FileStore/softia_mvp/knowledge"

INPUTS_BASE_DIR = f"{WORKSPACE_BASE_DIR}/inputs"
INDEXES_BASE_DIR = f"{WORKSPACE_BASE_DIR}/indexes"
OUTPUTS_BASE_DIR = f"{WORKSPACE_BASE_DIR}/outputs/generated_observations"

KNOWLEDGE_DIR = f"{KNOWLEDGE_BASE_DIR}/{METOD_ID}"
INPUT_DIR = f"{INPUTS_BASE_DIR}/{METOD_ID}"
INDEX_DIR = f"{INDEXES_BASE_DIR}/{METOD_ID}"
INDEX_FILE_NAME = "chunks_with_embeddings.json"

os.makedirs(INPUT_DIR, exist_ok=True)
os.makedirs(INDEX_DIR, exist_ok=True)
os.makedirs(OUTPUTS_BASE_DIR, exist_ok=True)

print("Knowledge folder:", KNOWLEDGE_DIR)
print("Input folder:", INPUT_DIR)
print("Index folder:", INDEX_DIR)
print("Output folder:", OUTPUTS_BASE_DIR)

# COMMAND ----------

# ============================================================
# CONSTANTS
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".docx",
    ".xlsx",
    ".html",
    ".htm",
    ".txt",
    ".md",
    ".py",
    ".sql",
    ".sas",
    ".png",
    ".jpg",
    ".jpeg",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
TEXT_EXTENSIONS = {".txt", ".md", ".py", ".sql", ".sas"}

DEFAULT_CHUNK_SIZE = 1800
DEFAULT_OVERLAP = 200
TOP_K_KNOWLEDGE = 6
TOP_K_EVIDENCE = 8

# COMMAND ----------

# ============================================================
# FILE DISCOVERY AND VALIDATION
# ============================================================


def list_supported_files(folder_path: str) -> List[str]:
    if not os.path.isdir(folder_path):
        print(f"WARNING: Folder does not exist while listing files: {folder_path}")
        return []

    supported_files = []
    for root, _, files in os.walk(folder_path):
        for file_name in sorted(files):
            file_path = os.path.join(root, file_name)
            extension = Path(file_name).suffix.lower()
            if extension in SUPPORTED_EXTENSIONS:
                supported_files.append(file_path)
            else:
                print(f"WARNING: Unsupported file skipped: {file_path}")

    return sorted(supported_files)


def validate_required_files(knowledge_files: List[str], input_files: List[str]) -> None:
    if not os.path.isdir(KNOWLEDGE_DIR):
        raise FileNotFoundError(f"Selected METOD knowledge folder does not exist: {KNOWLEDGE_DIR}")
    if not os.path.isdir(INPUT_DIR):
        raise FileNotFoundError(f"Selected METOD input folder does not exist: {INPUT_DIR}")
    if not knowledge_files:
        raise ValueError(f"No supported knowledge files found for {METOD_ID}: {KNOWLEDGE_DIR}")
    if not input_files:
        raise ValueError(f"No supported input evidence files found for {METOD_ID}: {INPUT_DIR}")


knowledge_files = list_supported_files(KNOWLEDGE_DIR)
input_files = list_supported_files(INPUT_DIR)
validate_required_files(knowledge_files, input_files)

print(f"Knowledge files discovered: {len(knowledge_files)}")
for file_path in knowledge_files:
    print(" -", file_path)

print(f"Input evidence files discovered: {len(input_files)}")
for file_path in input_files:
    print(" -", file_path)

# COMMAND ----------

# ============================================================
# PARSERS AND CHUNKING
# ============================================================


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\x00", "").strip()


def make_chunk_id(source_file: str, content_type: str, index: int) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_.-]+", "_", source_file)
    return f"{normalized}:{content_type}:{index}:{uuid.uuid4().hex[:8]}"


def split_text_into_chunks(
    text: str,
    source_file: str,
    content_type: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_OVERLAP,
    metadata: Optional[Dict[str, Any]] = None,
) -> List[Dict[str, Any]]:
    clean_text = re.sub(r"\n{3,}", "\n\n", safe_text(text))
    if not clean_text:
        return []
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")

    chunks = []
    start = 0
    chunk_index = 0
    while start < len(clean_text):
        end = min(start + chunk_size, len(clean_text))
        chunk_text = clean_text[start:end].strip()
        if chunk_text:
            chunk_metadata = dict(metadata or {})
            chunk_metadata.update(
                {
                    "source_file": source_file,
                    "file_path": chunk_metadata.get("file_path", ""),
                    "file_type": Path(source_file).suffix.lower(),
                    "content_type": content_type,
                    "char_start": start,
                    "char_end": end,
                }
            )
            chunks.append(
                {
                    "chunk_id": make_chunk_id(source_file, content_type, chunk_index),
                    "text": chunk_text,
                    "metadata": chunk_metadata,
                }
            )
            chunk_index += 1

        if end >= len(clean_text):
            break
        start = max(0, end - overlap)

    return chunks


def parse_docx(file_path: str) -> List[Dict[str, Any]]:
    source_file = os.path.basename(file_path)
    metadata = {"file_path": file_path}
    document = Document(file_path)

    chunks = []
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    paragraph_text = "\n\n".join(paragraphs)
    chunks.extend(split_text_into_chunks(paragraph_text, source_file, "docx_paragraphs", metadata=metadata))

    table_blocks = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [safe_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            table_blocks.append(f"Table {table_index}\n" + "\n".join(rows))

    table_text = "\n\n".join(table_blocks)
    chunks.extend(split_text_into_chunks(table_text, source_file, "docx_tables", metadata=metadata))
    return chunks


def parse_xlsx(file_path: str) -> List[Dict[str, Any]]:
    source_file = os.path.basename(file_path)
    workbook = load_workbook(file_path, data_only=True, read_only=True)
    chunks = []

    for worksheet in workbook.worksheets:
        buffered_rows = []
        row_start = None
        row_end = None
        current_size = 0

        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = [safe_text(value) for value in row]
            if not any(values):
                continue

            line = " | ".join(values)
            if row_start is None:
                row_start = row_number

            buffered_rows.append(f"Row {row_number}: {line}")
            row_end = row_number
            current_size += len(line)

            if current_size >= DEFAULT_CHUNK_SIZE:
                text = f"Sheet: {worksheet.title}\n" + "\n".join(buffered_rows)
                chunks.extend(
                    split_text_into_chunks(
                        text,
                        source_file,
                        "xlsx_rows",
                        metadata={
                            "file_path": file_path,
                            "sheet_name": worksheet.title,
                            "row_start": row_start,
                            "row_end": row_end,
                        },
                    )
                )
                buffered_rows = []
                row_start = None
                row_end = None
                current_size = 0

        if buffered_rows:
            text = f"Sheet: {worksheet.title}\n" + "\n".join(buffered_rows)
            chunks.extend(
                split_text_into_chunks(
                    text,
                    source_file,
                    "xlsx_rows",
                    metadata={
                        "file_path": file_path,
                        "sheet_name": worksheet.title,
                        "row_start": row_start,
                        "row_end": row_end,
                    },
                )
            )

    workbook.close()
    return chunks


def parse_html(file_path: str) -> List[Dict[str, Any]]:
    source_file = os.path.basename(file_path)
    with open(file_path, "r", encoding="utf-8", errors="replace") as file:
        html = file.read()

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    chunks = []
    headings = []
    for heading in soup.find_all(["h1", "h2", "h3", "h4"]):
        heading_text = heading.get_text(" ", strip=True)
        if heading_text:
            headings.append(heading_text)

    visible_text = soup.get_text("\n", strip=True)
    combined_text = "\n".join([f"Heading: {h}" for h in headings]) + "\n\n" + visible_text
    chunks.extend(
        split_text_into_chunks(
            combined_text,
            source_file,
            "html_visible_text",
            metadata={"file_path": file_path},
        )
    )

    table_blocks = []
    for table_index, table in enumerate(soup.find_all("table"), start=1):
        rows = []
        for row in table.find_all("tr"):
            cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            table_blocks.append(f"HTML table {table_index}\n" + "\n".join(rows))

    chunks.extend(
        split_text_into_chunks(
            "\n\n".join(table_blocks),
            source_file,
            "html_tables",
            metadata={"file_path": file_path},
        )
    )
    return chunks


def parse_text_or_script(file_path: str) -> List[Dict[str, Any]]:
    source_file = os.path.basename(file_path)
    extension = Path(file_path).suffix.lower().replace(".", "")
    with open(file_path, "r", encoding="utf-8", errors="replace") as file:
        text = file.read()

    return split_text_into_chunks(
        text,
        source_file,
        f"{extension}_content",
        metadata={"file_path": file_path},
    )


def describe_image(file_path: str) -> List[Dict[str, Any]]:
    source_file = os.path.basename(file_path)
    mime_type = mimetypes.guess_type(file_path)[0] or "image/png"

    with open(file_path, "rb") as image_file:
        image_base64 = base64.b64encode(image_file.read()).decode("utf-8")

    prompt = (
        "Describe this image as model validation evidence. Focus on visible tables, metrics, "
        "charts, warnings, output values, anomalies, and anything relevant to generating a "
        "validation observation. Do not invent values that are not visible."
    )

    response = azure_client.chat.completions.create(
        model=CHAT_DEPLOYMENT,
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime_type};base64,{image_base64}"},
                    },
                ],
            }
        ],
        temperature=0.0,
        max_tokens=800,
    )

    description = response.choices[0].message.content or ""
    return split_text_into_chunks(
        description,
        source_file,
        "image_description",
        metadata={"file_path": file_path, "mime_type": mime_type},
    )


def parse_file(file_path: str) -> List[Dict[str, Any]]:
    extension = Path(file_path).suffix.lower()

    try:
        if extension == ".docx":
            return parse_docx(file_path)
        if extension == ".xlsx":
            return parse_xlsx(file_path)
        if extension in {".html", ".htm"}:
            return parse_html(file_path)
        if extension in TEXT_EXTENSIONS:
            return parse_text_or_script(file_path)
        if extension in IMAGE_EXTENSIONS:
            return describe_image(file_path)

        print(f"WARNING: Unsupported file skipped: {file_path}")
        return []
    except Exception as error:
        print(f"WARNING: Failed to parse file: {file_path}. Reason: {error}")
        return []


def attach_chunk_runtime_fields(
    chunk: Dict[str, Any],
    metod_id: str,
    chunk_role: str,
) -> Dict[str, Any]:
    return {
        "chunk_id": chunk["chunk_id"],
        "metod_id": metod_id,
        "chunk_role": chunk_role,
        "text": chunk["text"],
        "metadata": chunk.get("metadata", {}),
        "embedding": [],
        "created_at": utc_now_iso(),
    }


def build_chunks_for_metod(
    metod_id: str,
    knowledge_files: List[str],
    input_files: List[str],
) -> List[Dict[str, Any]]:
    chunks = []

    for file_path in knowledge_files:
        parsed_chunks = parse_file(file_path)
        chunks.extend(
            attach_chunk_runtime_fields(chunk, metod_id, "knowledge")
            for chunk in parsed_chunks
        )

    for file_path in input_files:
        parsed_chunks = parse_file(file_path)
        chunks.extend(
            attach_chunk_runtime_fields(chunk, metod_id, "evidence")
            for chunk in parsed_chunks
        )

    if not any(chunk["chunk_role"] == "knowledge" for chunk in chunks):
        raise ValueError(f"No parseable knowledge chunks were created for {metod_id}.")
    if not any(chunk["chunk_role"] == "evidence" for chunk in chunks):
        raise ValueError(f"No parseable evidence chunks were created for {metod_id}.")

    return chunks

# COMMAND ----------

# ============================================================
# EMBEDDINGS, INDEXING, AND RETRIEVAL
# ============================================================


def get_embedding(text: str) -> List[float]:
    clean_text = safe_text(text)
    if not clean_text:
        raise ValueError("Cannot embed empty text.")

    response = azure_client.embeddings.create(
        model=EMBEDDING_DEPLOYMENT,
        input=clean_text,
    )
    return response.data[0].embedding


def add_embeddings_to_chunks(chunks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    embedded_chunks = []
    total = len(chunks)
    for index, chunk in enumerate(chunks, start=1):
        chunk["embedding"] = get_embedding(chunk["text"])
        embedded_chunks.append(chunk)
        if index % 10 == 0 or index == total:
            print(f"Embedded chunks: {index}/{total}")
    return embedded_chunks


def index_path_for_metod(metod_id: str) -> str:
    return f"{INDEXES_BASE_DIR}/{metod_id}/{INDEX_FILE_NAME}"


def save_chunks_index(metod_id: str, chunks: List[Dict[str, Any]]) -> str:
    metod_index_dir = f"{INDEXES_BASE_DIR}/{metod_id}"
    os.makedirs(metod_index_dir, exist_ok=True)
    index_path = index_path_for_metod(metod_id)
    with open(index_path, "w", encoding="utf-8") as file:
        json.dump(chunks, file, indent=2, ensure_ascii=False)
    return index_path


def load_chunks_index(metod_id: str) -> List[Dict[str, Any]]:
    index_path = index_path_for_metod(metod_id)
    if not os.path.isfile(index_path):
        raise FileNotFoundError(
            f"JSON index does not exist for {metod_id}: {index_path}. "
            "Run with rebuild_index=true first."
        )

    with open(index_path, "r", encoding="utf-8") as file:
        chunks = json.load(file)

    if not isinstance(chunks, list):
        raise ValueError(f"Invalid index format. Expected list of chunks: {index_path}")
    if not chunks:
        raise ValueError(f"Index is empty: {index_path}")
    return chunks


def cosine_similarity(a: List[float], b: List[float]) -> float:
    if not a or not b or len(a) != len(b):
        return 0.0

    dot_product = sum(x * y for x, y in zip(a, b))
    norm_a = math.sqrt(sum(x * x for x in a))
    norm_b = math.sqrt(sum(y * y for y in b))
    if norm_a == 0 or norm_b == 0:
        return 0.0
    return dot_product / (norm_a * norm_b)


def retrieve_chunks(
    query: str,
    chunks: List[Dict[str, Any]],
    chunk_role: str,
    top_k: int,
) -> List[Dict[str, Any]]:
    query_embedding = get_embedding(query)
    candidates = [chunk for chunk in chunks if chunk.get("chunk_role") == chunk_role]
    scored_chunks = []

    for chunk in candidates:
        score = cosine_similarity(query_embedding, chunk.get("embedding", []))
        scored_chunk = dict(chunk)
        scored_chunk["similarity_score"] = score
        scored_chunks.append(scored_chunk)

    scored_chunks.sort(key=lambda item: item["similarity_score"], reverse=True)
    return scored_chunks[:top_k]

# COMMAND ----------

# ============================================================
# CONTEXT AND PROMPT BUILDING
# ============================================================


def compact_chunk_for_prompt(chunk: Dict[str, Any]) -> Dict[str, Any]:
    metadata = chunk.get("metadata", {})
    return {
        "chunk_id": chunk.get("chunk_id"),
        "source_file": metadata.get("source_file"),
        "file_type": metadata.get("file_type"),
        "content_type": metadata.get("content_type"),
        "sheet_name": metadata.get("sheet_name"),
        "row_start": metadata.get("row_start"),
        "row_end": metadata.get("row_end"),
        "similarity_score": round(float(chunk.get("similarity_score", 0.0)), 4),
        "text": chunk.get("text", ""),
    }


def build_metod_context(
    metod_id: str,
    model_context: Dict[str, Any],
    chunks: List[Dict[str, Any]],
    top_k_knowledge: int = TOP_K_KNOWLEDGE,
    top_k_evidence: int = TOP_K_EVIDENCE,
) -> Dict[str, Any]:
    knowledge_query = (
        f"{metod_id} methodology rules thresholds criteria required analysis steps "
        "criticality rules action plan guidance missing information no observation conditions"
    )

    evidence_query = (
        "Relevant evidence from uploaded files to generate a validation observation "
        f"for {metod_id}. Model context: {json.dumps(model_context, ensure_ascii=False)}"
    )

    retrieved_knowledge = retrieve_chunks(knowledge_query, chunks, "knowledge", top_k_knowledge)
    retrieved_evidence = retrieve_chunks(evidence_query, chunks, "evidence", top_k_evidence)

    return {
        "metod_id": metod_id,
        "model_context": model_context,
        "retrieved_knowledge": retrieved_knowledge,
        "retrieved_evidence": retrieved_evidence,
    }


def build_observation_prompt(context: Dict[str, Any]) -> str:
    prompt_payload = {
        "metod_id": context["metod_id"],
        "model_context": context["model_context"],
        "retrieved_knowledge": [
            compact_chunk_for_prompt(chunk) for chunk in context["retrieved_knowledge"]
        ],
        "retrieved_evidence": [
            compact_chunk_for_prompt(chunk) for chunk in context["retrieved_evidence"]
        ],
    }

    return f"""
You are a METOD-specific model validation reviewer.

Generate one structured validation observation for the selected METOD.

Rules:
- Use METOD knowledge as the review criteria.
- Use runtime evidence as the facts to analyze.
- Use the widget model context as execution context.
- Do not invent facts, thresholds, metrics, sources, or conclusions.
- If evidence is insufficient, state what is missing.
- If no issue is supported by the evidence, say no observation is supported.
- Generate a specific observation, not a generic one.
- Justify criticality using the retrieved METOD knowledge.
- Recommend a concrete operational action plan.
- Cite source files used.
- Return valid JSON only. Do not wrap it in Markdown.

Required JSON structure:
{{
  "metod_id": "{context["metod_id"]}",
  "observation": "...",
  "criticality": "Low | Medium | High | Critical",
  "evidence": "...",
  "risk": "...",
  "action_plan": "...",
  "missing_information": ["..."],
  "confidence": 0.0,
  "sources": [
    {{
      "source_file": "...",
      "content_type": "...",
      "reason_used": "..."
    }}
  ]
}}

Context package:
{json.dumps(prompt_payload, indent=2, ensure_ascii=False)}
""".strip()


def print_retrieval_debug(context: Dict[str, Any], snippet_chars: int = 500) -> None:
    print("\n=== RAG retrieval debug: knowledge chunks ===")
    for index, chunk in enumerate(context["retrieved_knowledge"], start=1):
        metadata = chunk.get("metadata", {})
        snippet = chunk.get("text", "").replace("\n", " ")[:snippet_chars]
        print(
            f"{index}. score={chunk.get('similarity_score', 0):.4f} "
            f"source={metadata.get('source_file')} "
            f"content_type={metadata.get('content_type')}"
        )
        print(f"   snippet={snippet}")

    print("\n=== RAG retrieval debug: evidence chunks ===")
    for index, chunk in enumerate(context["retrieved_evidence"], start=1):
        metadata = chunk.get("metadata", {})
        snippet = chunk.get("text", "").replace("\n", " ")[:snippet_chars]
        print(
            f"{index}. score={chunk.get('similarity_score', 0):.4f} "
            f"source={metadata.get('source_file')} "
            f"content_type={metadata.get('content_type')}"
        )
        print(f"   snippet={snippet}")

# COMMAND ----------

# ============================================================
# OBSERVATION GENERATION AND OUTPUTS
# ============================================================


def parse_json_object(text: str) -> Dict[str, Any]:
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as first_error:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if not match:
            raise ValueError(f"Model returned invalid JSON: {first_error}\nRaw response:\n{text}")
        try:
            parsed = json.loads(match.group(0))
        except json.JSONDecodeError as second_error:
            raise ValueError(f"Model returned invalid JSON: {second_error}\nRaw response:\n{text}")

    if not isinstance(parsed, dict):
        raise ValueError("Model returned JSON, but the root value is not an object.")
    return parsed


def validate_observation_json(observation: Dict[str, Any]) -> None:
    required_keys = {
        "metod_id",
        "observation",
        "criticality",
        "evidence",
        "risk",
        "action_plan",
        "missing_information",
        "confidence",
        "sources",
    }
    missing_keys = sorted(required_keys - set(observation.keys()))
    if missing_keys:
        raise ValueError(f"Generated observation is missing required key(s): {missing_keys}")

    if observation["metod_id"] != METOD_ID:
        raise ValueError(
            f"Generated observation metod_id mismatch. Expected {METOD_ID}, "
            f"got {observation['metod_id']}"
        )

    if observation["criticality"] not in {"Low", "Medium", "High", "Critical"}:
        raise ValueError(
            "Generated observation has invalid criticality. "
            "Expected one of: Low, Medium, High, Critical."
        )

    if not isinstance(observation["missing_information"], list):
        raise ValueError("Generated observation missing_information must be a list.")
    if not isinstance(observation["sources"], list):
        raise ValueError("Generated observation sources must be a list.")


def generate_observation(context: Dict[str, Any]) -> Dict[str, Any]:
    prompt = build_observation_prompt(context)

    response = azure_client.chat.completions.create(
        model=CHAT_DEPLOYMENT,
        messages=[
            {
                "role": "system",
                "content": (
                    "You generate structured model-validation observations. "
                    "You return JSON only and never invent unsupported facts."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        temperature=0.0,
        max_tokens=1800,
        response_format={"type": "json_object"},
    )

    raw_content = response.choices[0].message.content or ""
    observation = parse_json_object(raw_content)
    validate_observation_json(observation)
    return observation


def observation_to_markdown(observation: Dict[str, Any]) -> str:
    missing_information = observation.get("missing_information") or []
    if missing_information:
        missing_md = "\n".join(f"- {item}" for item in missing_information)
    else:
        missing_md = "- None"

    sources = observation.get("sources") or []
    if sources:
        sources_md = "\n".join(
            "- source_file: {source_file}\n  content_type: {content_type}\n  reason_used: {reason_used}".format(
                source_file=source.get("source_file", ""),
                content_type=source.get("content_type", ""),
                reason_used=source.get("reason_used", ""),
            )
            for source in sources
        )
    else:
        sources_md = "- None"

    return f"""# Generated Observation

## METOD
{observation.get("metod_id", "")}

## Observation
{observation.get("observation", "")}

## Criticality
{observation.get("criticality", "")}

## Evidence
{observation.get("evidence", "")}

## Risk
{observation.get("risk", "")}

## Action Plan
{observation.get("action_plan", "")}

## Missing Information
{missing_md}

## Confidence
{observation.get("confidence", "")}

## Sources
{sources_md}
"""


def save_observation_outputs(observation: Dict[str, Any]) -> Dict[str, str]:
    os.makedirs(OUTPUTS_BASE_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = f"{observation['metod_id']}_observation_{timestamp}"
    json_path = f"{OUTPUTS_BASE_DIR}/{base_name}.json"
    markdown_path = f"{OUTPUTS_BASE_DIR}/{base_name}.md"

    with open(json_path, "w", encoding="utf-8") as json_file:
        json.dump(observation, json_file, indent=2, ensure_ascii=False)

    with open(markdown_path, "w", encoding="utf-8") as markdown_file:
        markdown_file.write(observation_to_markdown(observation))

    return {"json_path": json_path, "markdown_path": markdown_path}


def display_observation(observation: Dict[str, Any]) -> None:
    print(json.dumps(observation, indent=2, ensure_ascii=False))

    markdown = observation_to_markdown(observation)
    try:
        displayHTML(
            "<pre style='white-space: pre-wrap; font-family: inherit;'>"
            + markdown.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            + "</pre>"
        )
    except NameError:
        print(markdown)

# COMMAND ----------

# ============================================================
# NOTEBOOK EXECUTION FLOW
# ============================================================

if REBUILD_INDEX:
    print("Rebuilding JSON index from knowledge and evidence files.")
    chunks = build_chunks_for_metod(METOD_ID, knowledge_files, input_files)
    print(f"Chunks created: {len(chunks)}")
    chunks = add_embeddings_to_chunks(chunks)
    index_path = save_chunks_index(METOD_ID, chunks)
    print("Index saved:", index_path)
else:
    print("Loading existing JSON index.")
    chunks = load_chunks_index(METOD_ID)
    print(f"Chunks loaded: {len(chunks)}")

context = build_metod_context(
    METOD_ID,
    MODEL_CONTEXT,
    chunks,
    top_k_knowledge=TOP_K_KNOWLEDGE,
    top_k_evidence=TOP_K_EVIDENCE,
)

print("Retrieved knowledge chunks:", len(context["retrieved_knowledge"]))
print("Retrieved evidence chunks:", len(context["retrieved_evidence"]))
print_retrieval_debug(context)

observation = generate_observation(context)
output_paths = save_observation_outputs(observation)

print("Observation JSON saved:", output_paths["json_path"])
print("Observation Markdown saved:", output_paths["markdown_path"])

display_observation(observation)
