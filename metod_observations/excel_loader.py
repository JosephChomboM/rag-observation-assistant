from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .schemas import ControlSpec, PruebaSpec

FLAG_COLUMN = "flag_on_wf"

REQUIRED_COLUMNS = {
    "Detalle_prueba",
    "Detalle",
    "Prompt",
    "Limites",
    "Criticidad",
    "Excepciones",
    FLAG_COLUMN,
}

ENABLED_FLAG_VALUES = {"1", "true", "si", "s", "x", "yes", "y", "verdadero"}
DISABLED_FLAG_VALUES = {"", "0", "false", "no", "n", "falso"}


def normalize_metod_id(sheet_name: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "", sheet_name.upper())


def slugify(value: str, max_len: int = 70) -> str:
    clean = unicodedata.normalize("NFKD", value)
    clean = clean.encode("ascii", "ignore").decode("ascii")
    clean = re.sub(r"[^a-zA-Z0-9]+", "_", clean).strip("_").lower()
    return clean[:max_len].strip("_") or "sin_nombre"


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def is_row_enabled(value: Any, sheet_name: str = "", row_number: int = 0) -> bool:
    """Return True only when `flag_on_wf` marks the row as active.

    A disabled row is treated as if it did not exist in the Excel: no markdown,
    no manifest entry, no embedding and no observation.
    """
    if isinstance(value, bool):
        return value

    text = safe_text(value).lower()
    if text.endswith(".0"):
        text = text[:-2]
    if text in ENABLED_FLAG_VALUES:
        return True
    if text in DISABLED_FLAG_VALUES:
        return False
    raise ValueError(
        f"Invalid '{FLAG_COLUMN}' value {value!r} in sheet '{sheet_name}' row {row_number}. "
        "Use 1 to activate the row or 0 to ignore it."
    )


def _headers(ws) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = {safe_text(value): index for index, value in enumerate(row)}
    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        raise ValueError(
            f"Missing required Excel column(s) in sheet '{ws.title}': {sorted(missing)}"
        )
    return headers


def load_pruebas_from_excel(
    excel_path: str | Path,
    sheet_name: str | None = None,
) -> dict[str, list[PruebaSpec]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    selected_sheets = [sheet_name] if sheet_name else workbook.sheetnames
    result: dict[str, list[PruebaSpec]] = {}

    for selected in selected_sheets:
        if selected not in workbook.sheetnames:
            raise ValueError(f"Sheet '{selected}' does not exist in {excel_path}")
        ws = workbook[selected]
        headers = _headers(ws)
        metod_id = normalize_metod_id(ws.title)
        prueba_names: list[str] = []
        rows_by_prueba: dict[str, list[dict[str, Any]]] = {}
        current_prueba = ""

        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            raw_prueba = safe_text(row[headers["Detalle_prueba"]])
            if raw_prueba:
                current_prueba = raw_prueba
            if not current_prueba:
                continue

            if not is_row_enabled(row[headers[FLAG_COLUMN]], ws.title, row_number):
                continue

            detalle = safe_text(row[headers["Detalle"]])
            prompt = safe_text(row[headers["Prompt"]])
            limites = safe_text(row[headers["Limites"]])
            criticidad = safe_text(row[headers["Criticidad"]])
            excepciones = safe_text(row[headers["Excepciones"]])
            if not any([detalle, prompt, limites, criticidad, excepciones]):
                continue

            if current_prueba not in rows_by_prueba:
                prueba_names.append(current_prueba)
                rows_by_prueba[current_prueba] = []

            rows_by_prueba[current_prueba].append(
                {
                    "row_number": row_number,
                    "detalle": detalle,
                    "prompt": prompt,
                    "limites": limites,
                    "criticidad": criticidad,
                    "excepciones": excepciones,
                }
            )

        pruebas: list[PruebaSpec] = []
        for prueba_index, nombre_prueba in enumerate(prueba_names, start=1):
            prueba_id = f"{metod_id}_P{prueba_index:02d}"
            controls = []
            for control_index, row in enumerate(rows_by_prueba[nombre_prueba], start=1):
                controls.append(
                    ControlSpec(
                        metod_id=metod_id,
                        prueba_id=prueba_id,
                        control_id=f"{prueba_id}_C{control_index:02d}",
                        nombre_prueba=nombre_prueba,
                        detalle=row["detalle"],
                        prompt=row["prompt"],
                        limites=row["limites"],
                        criticidad=row["criticidad"],
                        excepciones=row["excepciones"],
                        row_number=row["row_number"],
                    )
                )
            if controls:
                pruebas.append(
                    PruebaSpec(
                        metod_id=metod_id,
                        prueba_id=prueba_id,
                        nombre_prueba=nombre_prueba,
                        controls=controls,
                        source_file=Path(excel_path).name,
                    )
                )

        result[metod_id] = pruebas

    return result
