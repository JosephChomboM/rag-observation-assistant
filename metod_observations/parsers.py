from __future__ import annotations

import csv
import html
import os
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .schemas import ParsedDocument

SUPPORTED_EVIDENCE_EXTENSIONS = {
    ".docx",
    ".xlsx",
    ".html",
    ".htm",
    ".txt",
    ".md",
    ".py",
    ".sql",
    ".sas",
    ".csv",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
}

TEXT_EXTENSIONS = {".txt", ".md", ".py", ".sql", ".sas"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}


def parse_text_file(path: str | Path) -> list[ParsedDocument]:
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    return [ParsedDocument(str(path), "text", text)]


def parse_csv_file(path: str | Path) -> list[ParsedDocument]:
    rows = []
    with Path(path).open("r", encoding="utf-8", errors="replace", newline="") as file:
        reader = csv.reader(file)
        for index, row in enumerate(reader, start=1):
            rows.append(" | ".join(str(value) for value in row))
            if index >= 1000:
                rows.append("[CSV truncado a 1000 filas para el MVP]")
                break
    return [ParsedDocument(str(path), "csv", "\n".join(rows))]


def parse_xlsx_file(path: str | Path) -> list[ParsedDocument]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    docs = []
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        lines = [f"Hoja: {sheet_name}"]
        for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                lines.append(" | ".join(values))
            if row_index >= 300:
                lines.append("[Hoja truncada a 300 filas para el MVP]")
                break
        docs.append(
            ParsedDocument(
                str(path),
                "xlsx",
                "\n".join(lines),
                {"sheet_name": sheet_name},
            )
        )
    return docs


def parse_docx_file(path: str | Path) -> list[ParsedDocument]:
    try:
        from docx import Document
    except ImportError as error:
        raise ImportError("python-docx is required to parse .docx evidence.") from error

    doc = Document(path)
    lines = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(doc.tables, start=1):
        lines.append(f"Tabla {table_index}")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [ParsedDocument(str(path), "docx", "\n".join(lines))]


def parse_html_file(path: str | Path) -> list[ParsedDocument]:
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        clean = html.unescape(raw)
        return [ParsedDocument(str(path), "html", clean)]
    soup = BeautifulSoup(raw, "html.parser")
    return [ParsedDocument(str(path), "html", soup.get_text("\n", strip=True))]


def parse_pdf_file(path: str | Path) -> list[ParsedDocument]:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ImportError("pypdf is required to parse .pdf evidence.") from error

    reader = PdfReader(str(path))
    docs = []
    for page_index, page in enumerate(reader.pages, start=1):
        docs.append(
            ParsedDocument(
                str(path),
                "pdf",
                page.extract_text() or "",
                {"page": page_index},
            )
        )
    return docs


def parse_image_file(path: str | Path) -> list[ParsedDocument]:
    text = (
        f"Imagen de evidencia: {Path(path).name}. "
        "El MVP registra la imagen como evidencia disponible; "
        "si se requiere lectura visual, usar un modelo con vision en una etapa posterior."
    )
    return [ParsedDocument(str(path), "image", text)]


def parse_evidence_file(path: str | Path) -> list[ParsedDocument]:
    suffix = Path(path).suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return parse_text_file(path)
    if suffix == ".csv":
        return parse_csv_file(path)
    if suffix == ".xlsx":
        return parse_xlsx_file(path)
    if suffix == ".docx":
        return parse_docx_file(path)
    if suffix in {".html", ".htm"}:
        return parse_html_file(path)
    if suffix == ".pdf":
        return parse_pdf_file(path)
    if suffix in IMAGE_EXTENSIONS:
        return parse_image_file(path)
    raise ValueError(f"Unsupported evidence extension: {path}")


def list_evidence_files(input_dir: str | Path) -> list[str]:
    base = Path(input_dir)
    if not base.exists():
        return []
    return sorted(
        str(path)
        for path in base.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EVIDENCE_EXTENSIONS
    )


def parse_evidence_files(paths: Iterable[str | Path]) -> list[ParsedDocument]:
    docs: list[ParsedDocument] = []
    for path in paths:
        docs.extend(parse_evidence_file(path))
    return docs
